"""Tests for build_docx.py. Skips if python-docx is not installed."""

import hashlib
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx.shared import Inches, Mm

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "build_docx.py"
FIXTURES = Path(__file__).resolve().parent / "fixtures"

# 1x1 transparent PNG
MIN_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
    "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
)


def load_module():
    import importlib.util

    spec = importlib.util.spec_from_file_location("build_docx_mod", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def mod():
    return load_module()


def copy_project(tmp_path, name):
    dst = tmp_path / name
    shutil.copytree(FIXTURES / name, dst)
    drawings = dst / "07-drawings"
    drawings.mkdir(parents=True, exist_ok=True)
    image = drawings / "fig1.png"
    image.write_bytes(MIN_PNG)
    digest = hashlib.sha256(MIN_PNG).hexdigest()
    fig_id = "图1" if name == "tiny-cn" else "FIG. 1"
    manifest = {
        "sheets": [
            {
                "id": fig_id,
                "image": "fig1.png",
                "source": "fig1.dot",
                "source_kind": "graphviz",
                "sha256": digest,
            }
        ]
    }
    (drawings / "fig1.dot").write_text("digraph G { a -> b }\n", encoding="utf-8")
    (drawings / "sheet_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return dst


def media_names(path):
    with zipfile.ZipFile(path) as zf:
        return [name for name in zf.namelist() if name.startswith("word/media/")]


def paragraph_text(path):
    document = docx.Document(str(path))
    return [p.text for p in document.paragraphs if p.text.strip()]


def assert_length_near(actual, expected, slack=5000):
    """OOXML stores page size in twips, so 210 mm round-trips a few EMU off Mm()."""
    assert abs(int(actual) - int(expected)) <= slack, (int(actual), int(expected))


def all_xml(path):
    with zipfile.ZipFile(path) as zf:
        return "\n".join(
            zf.read(name).decode("utf-8")
            for name in zf.namelist()
            if name.endswith(".xml")
        )


def test_cn_package(mod, tmp_path):
    project = copy_project(tmp_path, "tiny-cn")
    out = tmp_path / "out-cn"
    code = mod.main(["--jurisdiction", "cn", "--project", str(project), "--out", str(out)])
    assert code == 0
    expected = [
        "权利要求书.docx",
        "说明书.docx",
        "说明书摘要.docx",
        "说明书附图.docx",
        "摘要附图.docx",
    ]
    for name in expected:
        assert (out / name).is_file(), name

    spec = paragraph_text(out / "说明书.docx")
    numbered = [t for t in spec if t.startswith("[")]
    assert numbered[0].startswith("[0001]")
    assert numbered[1].startswith("[0002]")
    assert any("传感装置" in t for t in numbered)
    assert any("【待发明人确认】" in t for t in numbered)

    claims = "\n".join(paragraph_text(out / "权利要求书.docx"))
    assert "一种传感装置" in claims
    assert "[0001]" not in claims

    abstract = "\n".join(paragraph_text(out / "说明书摘要.docx"))
    assert "壳体" in abstract
    assert "[0001]" not in abstract

    drawings = "\n".join(paragraph_text(out / "说明书附图.docx"))
    assert "fig1.png" in drawings
    assert "图1" in drawings
    assert "figure_list.md" not in drawings
    assert media_names(out / "说明书附图.docx")

    brief = "\n".join(paragraph_text(out / "摘要附图.docx"))
    assert "fig1.png" in brief
    assert media_names(out / "摘要附图.docx")
    hashes = json.loads((out / "sheet_hashes.json").read_text(encoding="utf-8"))
    assert hashes["sheets"][0]["sha256"] == hashlib.sha256(MIN_PNG).hexdigest()

    document = docx.Document(str(out / "说明书.docx"))
    section = document.sections[0]
    assert_length_near(section.page_width, Mm(210))
    assert_length_near(section.page_height, Mm(297))
    xml = all_xml(out / "说明书.docx")
    assert "宋体" in xml
    assert "SimSun" in xml


def test_us_package(mod, tmp_path):
    project = copy_project(tmp_path, "tiny-us")
    out = tmp_path / "out-us"
    code = mod.main(["--jurisdiction", "us", "--project", str(project), "--out", str(out)])
    assert code == 0
    for name in ("specification.docx", "claims.docx", "abstract.docx"):
        assert (out / name).is_file()
    assert not (out / "说明书.docx").exists()

    spec = [t for t in paragraph_text(out / "specification.docx") if t.startswith("[")]
    assert spec[0].startswith("[0001]")
    assert spec[1].startswith("[0002]")
    assert any("sensor device" in t for t in spec)

    claims = "\n".join(paragraph_text(out / "claims.docx"))
    assert "comprising a housing" in claims
    assert "[0001]" not in claims

    document = docx.Document(str(out / "specification.docx"))
    section = document.sections[0]
    assert int(section.page_width) == int(Inches(8.5))
    assert int(section.page_height) == int(Inches(11))
    xml = all_xml(out / "specification.docx")
    assert "Times New Roman" in xml


def test_pct_uses_a4_and_english_filenames(mod, tmp_path):
    project = copy_project(tmp_path, "tiny-us")
    out = tmp_path / "out-pct"
    code = mod.main(["--jurisdiction", "pct", "--project", str(project), "--out", str(out)])
    assert code == 0
    assert (out / "specification.docx").is_file()
    document = docx.Document(str(out / "specification.docx"))
    section = document.sections[0]
    assert_length_near(section.page_width, Mm(210))
    assert_length_near(section.page_height, Mm(297))
    spec = "\n".join(paragraph_text(out / "specification.docx"))
    assert "[0001]" in spec


def test_png_changed_after_hash_exits_1_and_writes_nothing(mod, tmp_path):
    """Adversarial: the manifest is honest, then the PNG bytes change."""
    project = copy_project(tmp_path, "tiny-cn")
    image = project / "07-drawings" / "fig1.png"
    image.write_bytes(image.read_bytes() + b"\x00")
    out = tmp_path / "tampered-out"
    code = mod.main(["--jurisdiction", "cn", "--project", str(project), "--out", str(out)])
    assert code == 1
    assert not out.exists() or not list(out.glob("*.docx"))
    assert not (out / "sheet_hashes.json").exists()


def test_checksum_mismatch_exits_1_and_writes_nothing(mod, tmp_path):
    project = copy_project(tmp_path, "tiny-cn")
    manifest_path = project / "07-drawings" / "sheet_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["sheets"][0]["sha256"] = "0" * 64
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    out = tmp_path / "bad-out"
    code = mod.main(["--jurisdiction", "cn", "--project", str(project), "--out", str(out)])
    assert code == 1
    assert not out.exists() or not list(out.glob("*.docx"))


def test_forbidden_source_exits_1(mod, tmp_path):
    project = copy_project(tmp_path, "tiny-cn")
    manifest_path = project / "07-drawings" / "sheet_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["sheets"][0]["source_kind"] = "image_model"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    code = mod.main(["--jurisdiction", "cn", "--project", str(project), "--out", str(tmp_path / "forbid-out")])
    assert code == 1


def test_missing_claims_exits_2(mod, tmp_path):
    project = tmp_path / "empty"
    project.mkdir()
    code = mod.main(["--jurisdiction", "cn", "--project", str(project), "--out", str(tmp_path / "out")])
    assert code == 2


def test_cli_cn(tmp_path):
    project = copy_project(tmp_path, "tiny-cn")
    out = tmp_path / "cli-out"
    proc = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--jurisdiction",
            "cn",
            "--project",
            str(project),
            "--out",
            str(out),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode == 0, proc.stderr
    assert (out / "权利要求书.docx").is_file()
    assert (out / "说明书.docx").is_file()

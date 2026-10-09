#!/usr/bin/env python3
"""Build filing-oriented Word documents from a patent project directory.

    build_docx.py --jurisdiction cn|us|pct --project <dir> --out <dir>

CN (A4, SimSun / 宋体)
    权利要求书.docx
    说明书.docx                  paragraph numbers [0001]
    说明书摘要.docx
    说明书附图.docx              placeholder list of images in 07-drawings/
    摘要附图.docx                placeholder; uses patent.yaml abstract_figure

US (Letter, Times New Roman, 1.5 line spacing)
    specification.docx           paragraph numbers [0001]
    claims.docx
    abstract.docx

PCT
    Same English filenames as US, on A4 (PCT Rule 11). Body text is Times New Roman.
    This is a drafting package, not a RO/IB form set.

Inputs under --project
    patent.yaml                  optional title and abstract_figure
    04-claims/claims.md
    05-description/description.md
    06-abstract/abstract.md
    07-drawings/                 image files; markdown notes are not listed as figures

Exit codes
    0  documents written
    2  missing input, bad arguments, or python-docx is not installed
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Mm, Pt
except ImportError:  # pragma: no cover - tests skip when python-docx is absent
    Document = None

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".gif", ".bmp", ".emf", ".wmf", ".svg"}

CN_OUTPUTS = (
    "权利要求书.docx",
    "说明书.docx",
    "说明书摘要.docx",
    "说明书附图.docx",
    "摘要附图.docx",
)

US_OUTPUTS = (
    "specification.docx",
    "claims.docx",
    "abstract.docx",
)


def load_simple_yaml(path: Path) -> dict[str, str]:
    data: dict[str, str] = {}
    if not path.is_file():
        return data
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        data[key.strip()] = value.strip().strip("\"'")
    return data


def markdown_blocks(text: str) -> list[str]:
    """Blank-line paragraphs. Hash headings become plain title text. Numbered claim lines start a block."""
    blocks: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        if buf:
            blocks.append(" ".join(buf).strip())
            buf.clear()

    for line in text.splitlines():
        if not line.strip():
            flush()
            continue
        heading = re.match(r"^#{1,6}\s+(.*)$", line)
        if heading:
            flush()
            title = heading.group(1).strip()
            if title:
                blocks.append(title)
            continue
        if re.match(r"^\d+\s*[\.．]\s+", line) and buf:
            flush()
        buf.append(line.strip())
    flush()
    return [block for block in blocks if block]


def list_images(drawings: Path) -> list[Path]:
    if not drawings.is_dir():
        return []
    files = [path for path in drawings.iterdir() if path.is_file() and path.suffix.lower() in IMAGE_EXT]
    return sorted(files, key=lambda path: path.name.lower())


def _set_run_font(run, jurisdiction: str) -> None:
    if jurisdiction == "cn":
        ascii_name = "SimSun"
        east_asia = "宋体"
    else:
        ascii_name = "Times New Roman"
        east_asia = "Times New Roman"
    run.font.name = ascii_name
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")


def _apply_style(document, jurisdiction: str) -> None:
    style = document.styles["Normal"]
    if jurisdiction == "cn":
        style.font.name = "SimSun"
        style.font.size = Pt(12)
        style.element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
        style.element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    else:
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _configure_section(section, jurisdiction: str) -> None:
    if jurisdiction == "us":
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
    else:
        # CN practice and PCT Rule 11 use A4.
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)


def new_document(jurisdiction: str, title: str):
    document = Document()
    _apply_style(document, jurisdiction)
    _configure_section(document.sections[0], jurisdiction)
    document.core_properties.title = title
    document.core_properties.language = "zh-CN" if jurisdiction == "cn" else "en-US"
    return document


def add_paragraph(document, text: str, jurisdiction: str, number: int | None = None):
    paragraph = document.add_paragraph()
    if jurisdiction in ("us", "pct"):
        paragraph.paragraph_format.line_spacing = 1.5
    shown = f"[{number:04d}] {text}" if number is not None else text
    run = paragraph.add_run(shown)
    _set_run_font(run, jurisdiction)
    return paragraph


def add_blocks(document, blocks: list[str], jurisdiction: str, numbered: bool) -> None:
    for index, block in enumerate(blocks, start=1):
        add_paragraph(document, block, jurisdiction, index if numbered else None)


def _read(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def build_package(project: Path, out_dir: Path, jurisdiction: str) -> list[Path]:
    if Document is None:
        raise RuntimeError("python-docx is not installed. Install it with: python3 -m pip install --user python-docx")
    claims_path = project / "04-claims" / "claims.md"
    desc_path = project / "05-description" / "description.md"
    abstract_path = project / "06-abstract" / "abstract.md"
    missing = [str(path.relative_to(project)) for path in (claims_path, desc_path, abstract_path) if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required project files: " + ", ".join(missing))

    meta = load_simple_yaml(project / "patent.yaml")
    title = meta.get("title", "")
    claims = markdown_blocks(_read(claims_path))
    description = markdown_blocks(_read(desc_path))
    abstract = markdown_blocks(_read(abstract_path))
    images = list_images(project / "07-drawings")
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    if jurisdiction == "cn":
        written.append(_save_cn_claims(out_dir, jurisdiction, title, claims))
        written.append(_save_numbered(out_dir / "说明书.docx", jurisdiction, title or "说明书", description))
        written.append(_save_plain(out_dir / "说明书摘要.docx", jurisdiction, "说明书摘要", title, abstract))
        written.append(_save_cn_drawings(out_dir, jurisdiction, images))
        written.append(_save_cn_abstract_figure(out_dir, jurisdiction, images, meta.get("abstract_figure", "")))
    else:
        written.append(_save_numbered(out_dir / "specification.docx", jurisdiction, title or "Specification", description))
        written.append(_save_plain(out_dir / "claims.docx", jurisdiction, "Claims", title, claims))
        written.append(_save_plain(out_dir / "abstract.docx", jurisdiction, "Abstract", title, abstract))
    return written


def _save_numbered(path: Path, jurisdiction: str, title: str, blocks: list[str]) -> Path:
    document = new_document(jurisdiction, title)
    if title:
        add_paragraph(document, title, jurisdiction, None)
    add_blocks(document, blocks, jurisdiction, numbered=True)
    document.save(path)
    return path


def _save_plain(path: Path, jurisdiction: str, heading: str, title: str, blocks: list[str]) -> Path:
    document = new_document(jurisdiction, title or heading)
    add_paragraph(document, heading, jurisdiction, None)
    if title and title != heading:
        add_paragraph(document, title, jurisdiction, None)
    add_blocks(document, blocks, jurisdiction, numbered=False)
    document.save(path)
    return path


def _save_cn_claims(out_dir: Path, jurisdiction: str, title: str, claims: list[str]) -> Path:
    return _save_plain(out_dir / "权利要求书.docx", jurisdiction, "权利要求书", title, claims)


def _save_cn_drawings(out_dir: Path, jurisdiction: str, images: list[Path]) -> Path:
    document = new_document(jurisdiction, "说明书附图")
    add_paragraph(document, "说明书附图", jurisdiction, None)
    if not images:
        add_paragraph(
            document,
            "【待发明人确认】07-drawings/ 中尚未放入附图文件。",
            jurisdiction,
            None,
        )
    else:
        add_paragraph(
            document,
            "以下为附图占位清单。正式图样以 07-drawings/ 中的图像文件为准。",
            jurisdiction,
            None,
        )
        for index, image in enumerate(images, start=1):
            add_paragraph(document, f"图{index}    {image.name}", jurisdiction, None)
    path = out_dir / "说明书附图.docx"
    document.save(path)
    return path


def _save_cn_abstract_figure(out_dir: Path, jurisdiction: str, images: list[Path], chosen: str) -> Path:
    document = new_document(jurisdiction, "摘要附图")
    add_paragraph(document, "摘要附图", jurisdiction, None)
    names = {image.name for image in images}
    if chosen:
        add_paragraph(document, f"指定摘要附图：{chosen}", jurisdiction, None)
        if chosen not in names:
            add_paragraph(
                document,
                "【待发明人确认】未在 07-drawings/ 中找到该文件。",
                jurisdiction,
                None,
            )
    else:
        add_paragraph(
            document,
            "【待发明人确认】请从下列附图中选择一幅作为摘要附图。",
            jurisdiction,
            None,
        )
    if images:
        add_paragraph(document, "可用附图：", jurisdiction, None)
        for image in images:
            add_paragraph(document, image.name, jurisdiction, None)
    else:
        add_paragraph(document, "【待发明人确认】07-drawings/ 中尚未放入附图文件。", jurisdiction, None)
    path = out_dir / "摘要附图.docx"
    document.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export a patent project to jurisdiction-specific .docx files.")
    parser.add_argument("--jurisdiction", required=True, choices=["cn", "us", "pct"])
    parser.add_argument("--project", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args(argv)
    if Document is None:
        print(
            "python-docx is not installed. Install it with: python3 -m pip install --user python-docx",
            file=sys.stderr,
        )
        return 2
    if not args.project.is_dir():
        print(f"Project directory not found: {args.project}", file=sys.stderr)
        return 2
    try:
        written = build_package(args.project, args.out, args.jurisdiction)
    except FileNotFoundError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    except RuntimeError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print(f"Jurisdiction: {args.jurisdiction}")
    print("Wrote:")
    for path in written:
        print(f"- {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
